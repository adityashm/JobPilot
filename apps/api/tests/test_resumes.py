import io
import docx
import pytest
from fastapi.testclient import TestClient
from app.services.resume_parser import extract_text_from_file, clean_text


def create_sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("Aditya Sharma - Software Engineer", level=1)
    doc.add_paragraph("Email: aditya@example.com | Phone: (555) 123-4567")
    doc.add_paragraph("LinkedIn: https://linkedin.com/in/adityashm | GitHub: https://github.com/adityashm")
    doc.add_paragraph("Experienced backend developer with 3+ years of experience building Python and FastAPI microservices.")
    doc.add_paragraph("Skills: Python, FastAPI, Docker, PostgreSQL, React, Next.js, Redis, PyTest")
    
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_resume_parser_docx_and_clean():
    docx_bytes = create_sample_docx_bytes()
    text = extract_text_from_file(docx_bytes, "resume.docx")
    assert "Aditya Sharma" in text
    assert "FastAPI" in text
    assert "3+ years" in text


def test_resume_parser_txt():
    content = b"Candidate Name\n\nSkills: Python, Go, Docker, Kubernetes\nYears of experience: 4 years"
    text = extract_text_from_file(content, "resume.txt")
    assert "Candidate Name" in text
    assert "Kubernetes" in text


def test_resume_upload_and_profile_sync(client: TestClient):
    # 1. Register user
    reg_resp = client.post(
        "/api/v1/auth/register",
        json={"email": "resume_user@example.com", "password": "password123", "full_name": "Resume User"},
    )
    assert reg_resp.status_code == 201
    token = reg_resp.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # 2. Upload DOCX resume
    docx_bytes = create_sample_docx_bytes()
    files = {"file": ("aditya_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    data = {"auto_update_profile": "true"}

    upload_resp = client.post("/api/v1/resumes", headers=headers, files=files, data=data)
    assert upload_resp.status_code == 200, upload_resp.text
    res_data = upload_resp.json()
    assert res_data["resume"]["filename"] == "aditya_resume.docx"
    assert res_data["resume"]["is_primary"] is True
    assert "Python" in res_data["extracted_profile"]["skills"]
    assert "FastAPI" in res_data["extracted_profile"]["skills"]
    assert res_data["extracted_profile"]["experience_years"] >= 3.0

    # 3. Verify user's profile was automatically updated
    profile_resp = client.get("/api/v1/profile", headers=headers)
    assert profile_resp.status_code == 200
    prof = profile_resp.json()
    assert "Python" in prof["skills"]
    assert "FastAPI" in prof["skills"]
    assert prof["linkedin_url"] == "https://linkedin.com/in/adityashm"
    assert prof["github_url"] == "https://github.com/adityashm"

    # 4. List resumes
    list_resp = client.get("/api/v1/resumes", headers=headers)
    assert list_resp.status_code == 200
    resumes_list = list_resp.json()
    assert len(resumes_list) == 1
    resume_id = resumes_list[0]["id"]

    # 5. Get resume detail
    detail_resp = client.get(f"/api/v1/resumes/{resume_id}", headers=headers)
    assert detail_resp.status_code == 200
    assert "Aditya Sharma" in detail_resp.json()["raw_text"]

    # 6. Upload a second resume (txt)
    txt_bytes = b"Jane Doe - Frontend Engineer\nSkills: React, TypeScript, Tailwind CSS\n2 years of experience"
    files2 = {"file": ("resume_frontend.txt", txt_bytes, "text/plain")}
    upload2 = client.post("/api/v1/resumes", headers=headers, files=files2, data={"auto_update_profile": "false"})
    assert upload2.status_code == 200
    res2_id = upload2.json()["resume"]["id"]
    assert upload2.json()["resume"]["is_primary"] is False

    # 7. Set second resume as primary
    primary_resp = client.post(f"/api/v1/resumes/{res2_id}/set-primary", headers=headers)
    assert primary_resp.status_code == 200
    assert primary_resp.json()["is_primary"] is True

    # Check first is no longer primary
    detail1 = client.get(f"/api/v1/resumes/{resume_id}", headers=headers).json()
    assert detail1["is_primary"] is False

    # 8. Delete resume
    del_resp = client.delete(f"/api/v1/resumes/{resume_id}", headers=headers)
    assert del_resp.status_code == 204
    list_after = client.get("/api/v1/resumes", headers=headers).json()
    assert len(list_after) == 1


def test_resume_upload_validation(client: TestClient):
    reg = client.post("/api/v1/auth/register", json={"email": "val@example.com", "password": "password123"})
    token = reg.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    # Invalid extension
    bad_file = {"file": ("malicious.exe", b"binarycontent", "application/octet-stream")}
    bad_resp = client.post("/api/v1/resumes", headers=headers, files=bad_file)
    assert bad_resp.status_code == 400

    # Empty file
    empty_file = {"file": ("empty.txt", b"", "text/plain")}
    empty_resp = client.post("/api/v1/resumes", headers=headers, files=empty_file)
    assert empty_resp.status_code == 400


def test_comprehensive_resume_details_extraction(client: TestClient):
    reg = client.post("/api/v1/auth/register", json={"email": "aditya_detailed@example.com", "password": "password123", "full_name": "Aditya Sharma"})
    token = reg.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    sample_resume = b"""
Aditya Sharma
Ghaziabad, Delhi NCR, India | +91 9876543210 | aditya@example.com
LinkedIn: https://linkedin.com/in/adityashm | GitHub: https://github.com/adityashm | Portfolio: https://adityashm.tech

PROFESSIONAL SUMMARY
Final-year B.Tech CSE student at IMS Engineering College with deep expertise in Python, FastAPI, React, PostgreSQL, Docker, and browser automation.

EDUCATION
IMS Engineering College, Ghaziabad (AKTU)
Bachelor of Technology in Computer Science and Engineering | CGPA: 8.31 / 10 | 2023 - 2027

TECHNICAL SKILLS
Languages: Python, Java, Go, JavaScript, TypeScript, Kotlin, SQL
Frontend: React, Next.js, Three.js, GSAP, Tailwind CSS, PyWebView
Backend: FastAPI, PostgreSQL, SQLite, Docker, Redis, REST APIs
Machine Learning & Mobile: Scikit-learn, Isolation Forest, Android (Jetpack Compose, Room)

WORK EXPERIENCE
Postix (postix.in) - Co-Founder & Full Stack Lead
June 2024 - Present
- Built AI marketing OS and attribution pixel architecture with vanilla JS landing page.
- Engineered companion native Android application using Kotlin, Jetpack Compose, and Room.

PROJECTS
AgentShield AI | Python, FastAPI, React, Isolation Forest | https://github.com/adityashm/AgentShield
- Security monitor for AI agents using Isolation Forest anomaly detection with 96% accuracy.
- Built dark-mode Command Center React dashboard with real-time audit telemetry.

CA Audit Assistant | React, FastAPI, PyWebView, SQLite | https://github.com/adityashm/ca-audit
- Desktop app automating audit engagements for Chartered Accountants (CARO 2020, GST reconciliation).
- 161 automated unit and integration tests bundled via PyInstaller for Windows.
"""

    files = {"file": ("aditya_detailed.txt", sample_resume, "text/plain")}
    upload_resp = client.post("/api/v1/resumes", headers=headers, files=files, data={"auto_update_profile": "true"})
    assert upload_resp.status_code == 200, upload_resp.text

    profile_resp = client.get("/api/v1/profile", headers=headers)
    assert profile_resp.status_code == 200
    prof = profile_resp.json()

    # Verify comprehensive details extraction and sync
    assert prof["phone"] == "+91 9876543210"
    assert "Ghaziabad" in (prof["location"] or "")
    assert prof["linkedin_url"] == "https://linkedin.com/in/adityashm"
    assert prof["github_url"] == "https://github.com/adityashm"
    assert prof["portfolio_url"] == "https://adityashm.tech"
    assert prof["experience_years"] >= 1.0
    assert len(prof["skills"]) >= 20
    assert "FastAPI" in prof["skills"]
    assert "Scikit-Learn" in prof["skills"] or "Scikit-learn" in prof["skills"]
    assert "Docker" in prof["skills"]
    assert "Kotlin" in prof["skills"]

    # Verify education, experience, and projects sections
    assert len(prof["education"]) >= 1
    assert "IMS Engineering College" in prof["education"][0]["institution"]
    assert "Bachelor of Technology" in prof["education"][0]["degree"]
    assert prof["education"][0]["start_year"] == 2023
    assert prof["education"][0]["end_year"] == 2027

    assert len(prof["experience"]) >= 1
    assert "Postix" in prof["experience"][0]["company"]
    assert "Co-Founder & Full Stack Lead" in prof["experience"][0]["role"]

    assert len(prof["projects"]) >= 2
    project_names = [p["name"] for p in prof["projects"]]
    assert any("AgentShield" in name for name in project_names)
    assert any("CA Audit" in name for name in project_names)

